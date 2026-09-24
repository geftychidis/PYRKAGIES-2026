#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Εισάγει στο LINKS τους συνδέσμους του βιβλίου στατιστικών ΠΥΡΚΑΓΙΕΣ 2026.

Το βιβλίο στατιστικών (φύλλο «Ελλάδα») κρατά ανά περιστατικό μια στήλη «Πηγές»
με πολλαπλά URL. Το script τα μετατρέπει σε γραμμές του LINKS, στη σωστή μηνιαία
ενότητα, παραλείποντας όσα URL υπάρχουν ήδη.

ΣΗΜΑΝΤΙΚΟ ΓΙΑ ΤΗΝ ΑΚΡΙΒΕΙΑ: το βιβλίο ΔΕΝ κρατά τίτλο άρθρου ανά URL. Στη στήλη
«Τίτλος» μπαίνει το όνομα του περιστατικού και κάθε τέτοια γραμμή σημειώνεται με
«(στατιστικά ΠΥΡΚΑΓΙΕΣ 2026)», ώστε να μην εμφανίζεται ως πραγματική κεφαλίδα.

Χρήση:
    python3 scripts/import_from_stats_xlsx.py                      # dry-run
    python3 scripts/import_from_stats_xlsx.py --apply
    python3 scripts/import_from_stats_xlsx.py --apply --since 2026-08-14
"""

import argparse
import collections
import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "ΠΥΡΚΑΓΙΕΣ_2026_ΣΤΑΤΙΣΤΙΚΑ.xlsx"
LINKS = ROOT / "LINKS_ΔΑΣΙΚΕΣ_ΠΥΡΚΑΓΙΕΣ_2026.md"
TAG = "(στατιστικά ΠΥΡΚΑΓΙΕΣ 2026)"

MONTHS = ["", "Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
          "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος"]

# Στήλες του φύλλου «Ελλάδα»
C_DATE, C_NAME, C_PERIF, C_NOMOS, C_DIMOS = 1, 5, 6, 7, 8
C_HA, C_STATUS, C_CAUSE, C_112, C_SRC, C_NOTES = 16, 19, 20, 21, 22, 23

SOURCES = {
    "ertnews.gr": "ΕΡΤ News", "tanea.gr": "Τα Νέα", "skai.gr": "ΣΚΑΪ",
    "newsit.gr": "Newsit", "naftemporiki.gr": "Ναυτεμπορική", "cnn.gr": "CNN Greece",
    "documentonews.gr": "Documento", "newsbomb.gr": "Newsbomb", "protothema.gr": "Πρώτο Θέμα",
    "kathimerini.gr": "Καθημερινή", "in.gr": "in.gr", "iefimerida.gr": "iefimerida",
    "news247.gr": "News247", "efsyn.gr": "Εφ.Συν.", "tovima.com": "Το Βήμα",
    "tovima.gr": "Το Βήμα", "real.gr": "Real.gr", "megatv.com": "MEGA",
    "ethnos.gr": "Έθνος", "thetoc.gr": "The TOC", "zougla.gr": "Ζούγκλα",
    "fireservice.gr": "Πυροσβεστικό Σώμα", "civilprotection.gov.gr": "Πολιτική Προστασία",
    "meteo.gr": "meteo.gr", "reuters.com": "Reuters", "apnews.com": "AP",
    "euronews.com": "Euronews", "bbc.com": "BBC", "aljazeera.com": "Al Jazeera",
    "ekathimerini.com": "eKathimerini", "firefightingreece.gr": "Fire Fighting Greece",
    "lifo.gr": "LiFO", "star.gr": "STAR", "alphatv.gr": "ALPHA", "ant1news.gr": "ANT1",
    "newpost.gr": "NewPost", "dikaiologitika.gr": "Δικαιολογητικά Νέα",
    "typosthes.gr": "Τύπος Θεσσαλονίκης", "voria.gr": "Voria", "cretalive.gr": "CretaLive",
}


def domain(url):
    return re.sub(r"^https?://(www\.)?", "", url).split("/")[0].lower()


def source_name(url):
    d = domain(url)
    if d in SOURCES:
        return SOURCES[d]
    core = d.split(":")[0]
    for k, v in SOURCES.items():
        if core.endswith("." + k):
            return v
    return core


def cell(ws, r, c):
    v = ws.cell(row=r, column=c).value
    return "" if v is None else str(v).strip()


def area_of(ws, r):
    perif, nomos, dimos = cell(ws, r, C_PERIF), cell(ws, r, C_NOMOS), cell(ws, r, C_DIMOS)
    head = perif or nomos or "Ελλάδα"
    tail = dimos or (nomos if perif and nomos and nomos not in head else "")
    return f"{head} / {tail}" if tail and tail not in head else head


def description(ws, r):
    bits = []
    ha = ws.cell(row=r, column=C_HA).value
    if isinstance(ha, (int, float)) and ha > 0:
        bits.append(f"{ha:g} ha")
    for c in (C_STATUS, C_CAUSE):
        v = cell(ws, r, c)
        if v and v not in ("—", "-"):
            bits.append(v)
    e = cell(ws, r, C_112)
    if e:
        first = re.split(r"[|·]", e)[0].strip()
        first = re.sub(r"^112\s*", "", first)          # χωρίς διπλό «112»
        if first:
            bits.append("112: " + first)
    if not bits:
        note = cell(ws, r, C_NOTES)
        if note:
            bits.append(re.split(r"[|·.]", note)[0].strip())
    desc = "· ".join(b for b in bits if b)
    if len(desc) > 130:
        desc = desc[:127].rstrip() + "…"
    return (desc + " " + TAG).strip() if desc else TAG


def clean(s):
    return s.replace("|", "/").replace("\n", " ").strip()


def collect(since):
    ws = load_workbook(XLSX, data_only=True)["Ελλάδα"]
    seen, out = set(), []
    for r in range(2, ws.max_row + 1):
        d = ws.cell(row=r, column=C_DATE).value
        d = d.date() if isinstance(d, dt.datetime) else d
        if not isinstance(d, dt.date) or d <= since:
            continue
        name = clean(cell(ws, r, C_NAME)) or "Πυρκαγιά"
        area, desc = clean(area_of(ws, r)), clean(description(ws, r))
        for u in re.findall(r"https?://[^\s;,|\)]+", cell(ws, r, C_SRC)):
            u = u.rstrip(".,;")
            if u in seen:
                continue
            seen.add(u)
            out.append((d, area, source_name(u), name, u, desc))
    return out


def month_key(line):
    m = re.match(r"\|\s*(\d{4})-(\d{2})-(\d{2})\s*\|", line)
    return (int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else (9999, 99, 99)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--since", default="2026-08-14", help="εισαγωγή περιστατικών ΜΕΤΑ από αυτή την ημερομηνία")
    ap.add_argument("--apply", action="store_true", help="χωρίς αυτό γίνεται μόνο dry-run")
    a = ap.parse_args()
    since = dt.date.fromisoformat(a.since)

    md = LINKS.read_text(encoding="utf-8")
    existing = set(re.findall(r"https?://[^\s|\)]+", md))
    rows = [r for r in collect(since) if r[4] not in existing]
    if not rows:
        print("Καμία νέα εγγραφή.")
        return

    by_month = collections.defaultdict(list)
    for d, area, src, title, url, desc in rows:
        by_month[(d.year, d.month)].append(
            f"| {d.isoformat()} | {area} | {src} | {title} | {url} | {desc} |")

    print(f"Νέοι σύνδεσμοι: {len(rows)}")
    for k in sorted(by_month):
        print(f"  {MONTHS[k[1]]} {k[0]}: {len(by_month[k])}")
    if not a.apply:
        print("\n(dry-run — τρέξε με --apply)")
        for line in by_month[sorted(by_month)[0]][:3]:
            print("  δείγμα:", line[:190])
        return

    lines = md.split("\n")
    # Ο Αύγουστος δεν είναι πια μερικός μήνας.
    lines = [l.replace("## Αύγουστος 2026 (1–14/8)", "## Αύγουστος 2026") for l in lines]
    md = "\n".join(lines)
    md = md.replace("- [Αύγουστος 2026 (1–14/8)](#αύγουστος-2026-1148)",
                    "- [Αύγουστος 2026](#αύγουστος-2026)")

    for (y, m) in sorted(by_month):
        head = f"## {MONTHS[m]} {y}"
        if head not in md:
            # νέα ενότητα μήνα, πριν από τις μόνιμες πηγές
            anchor = MONTHS[m].lower().replace(" ", "-")
            md = md.replace("- [Επίσημες & μόνιμες / ζωντανές πηγές]",
                            f"- [{MONTHS[m]} {y}](#{anchor}-{y})\n- [Επίσημες & μόνιμες / ζωντανές πηγές]", 1)
            block = (f"{head}\n\n**Σύνοψη μήνα:** Εγγραφές από το βιβλίο στατιστικών "
                     f"ΠΥΡΚΑΓΙΕΣ 2026.\n\n"
                     "| Ημ/νία | Περιοχή | Πηγή | Τίτλος | Σύνδεσμος | Περιγραφή |\n"
                     "|---|---|---|---|---|---|\n\n---\n\n")
            md = md.replace("## Επίσημες & μόνιμες / ζωντανές πηγές", block +
                            "## Επίσημες & μόνιμες / ζωντανές πηγές", 1)

        # βρες τον πίνακα της ενότητας και συγχώνευσε ταξινομημένα
        i = md.index(head)
        j = md.find("\n## ", i + 1)
        j = len(md) if j == -1 else j
        sec = md[i:j]
        body = sec.split("\n")
        start = next(k for k, l in enumerate(body) if l.startswith("|---"))
        end = start + 1
        while end < len(body) and body[end].startswith("|"):
            end += 1
        merged = sorted(body[start + 1:end] + by_month[(y, m)], key=month_key)
        md = md[:i] + "\n".join(body[:start + 1] + merged + body[end:]) + md[j:]

    total = len(set(re.findall(r"https?://[^\s|\)]+", md)))
    today = dt.date.today()
    md = re.sub(r"> \*\*Τελευταία ενημέρωση:\*\* \d{4}-\d{2}-\d{2}",
                f"> **Τελευταία ενημέρωση:** {today.isoformat()}", md)
    md = re.sub(r"> \*\*Σύνολο συνδέσμων:\*\*.*",
                f"> **Σύνολο συνδέσμων:** {total} μοναδικά URL", md)
    md = re.sub(r"> \*\*Περίοδος κάλυψης:\*\* 1 Μαΐου 2026 → [^\n]*",
                f"> **Περίοδος κάλυψης:** 1 Μαΐου 2026 → {today.strftime('%d/%m/%Y')}", md)
    md = md.replace("\n\n---\n\n*Αυτόματη συλλογή",
                    f"\n| {today.isoformat()} | +{len(rows)} (στατιστικά) | Αναδρομική εισαγωγή "
                    f"συνδέσμων από το βιβλίο ΠΥΡΚΑΓΙΕΣ_2026_ΣΤΑΤΙΣΤΙΚΑ.xlsx για την περίοδο "
                    f"{(since + dt.timedelta(days=1)).strftime('%d/%m')}–{today.strftime('%d/%m')}, "
                    f"που έλειπε λόγω βλάβης της καθημερινής routine. Η στήλη «Τίτλος» φέρει το "
                    f"όνομα του περιστατικού, όχι την κεφαλίδα του άρθρου. |\n\n---\n\n*Αυτόματη συλλογή", 1)
    LINKS.write_text(md, encoding="utf-8")
    print(f"\nOK: +{len(rows)} γραμμές· σύνολο μοναδικών URL στο LINKS: {total}")


if __name__ == "__main__":
    main()
