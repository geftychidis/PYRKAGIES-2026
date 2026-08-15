#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Αναδημιουργεί το ΠΡΟΟΔΟΣ_ΣΥΛΛΟΓΗΣ_2026.xlsx από το LINKS_ΔΑΣΙΚΕΣ_ΠΥΡΚΑΓΙΕΣ_2026.md.

Παράγει μία γραμμή ανά ημερολογιακή ημέρα (1/5/2026 → σήμερα ή τελευταία ημέρα με
δεδομένα), με τον αριθμό συνδέσμων που φέρουν εκείνη την ημερομηνία, το σωρευτικό
σύνολο, και δύο γραφήματα (ημερήσια κατανομή + σωρευτική καμπύλη).

Χρήση:  python3 scripts/build_progress_xlsx.py [YYYY-MM-DD]
        (προαιρετικό όρισμα: τελευταία ημέρα του πίνακα· default = σήμερα)
"""

import collections
import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
LINKS = ROOT / "LINKS_ΔΑΣΙΚΕΣ_ΠΥΡΚΑΓΙΕΣ_2026.md"
OUT = ROOT / "ΠΡΟΟΔΟΣ_ΣΥΛΛΟΓΗΣ_2026.xlsx"
SEASON_START = dt.date(2026, 5, 1)

MONTHS = ["Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
          "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος"]

HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
BASE_FILL = PatternFill("solid", fgColor="EAEFF2")
THIN = Side(style="thin", color="B7C3C9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def parse_links():
    """Επιστρέφει (counter ανά ημερομηνία, πλήθος μόνιμων/αχρονολόγητων πηγών)."""
    per_day = collections.Counter()
    undated = 0
    for line in LINKS.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|") or "http" not in line:
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 6:
            continue
        first = cells[0]
        if first in ("Ημ/νία", "Ημερομηνία ενημέρωσης") or set(first) <= set("-: "):
            continue
        m = re.match(r"(\d{4})-(\d{2})-(\d{2})", first)
        if m:
            per_day[dt.date(*map(int, m.groups()))] += 1
        else:
            undated += 1
    return per_day, undated


def build(per_day, undated, last_day):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ημερήσια Πρόοδος"

    headers = ["Ημερομηνία (μέρα/μήνας)", "Νέα Links ημέρας", "Σύνολο Links"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER

    # Γραμμή βάσης: μόνιμες / ζωντανές πηγές χωρίς ημερομηνία περιστατικού.
    running = undated
    ws.append(["— (μόνιμες)", undated, running])
    for c in ws[2]:
        c.fill = BASE_FILL
        c.border = BORDER
    ws.cell(row=2, column=1).comment = None

    first_data_row = 3
    day = SEASON_START
    while day <= last_day:
        n = per_day.get(day, 0)
        running += n
        ws.append([day.strftime("%d/%m"), n, running])
        day += dt.timedelta(days=1)
    last_data_row = ws.max_row

    for row in ws.iter_rows(min_row=first_data_row, max_row=last_data_row, max_col=3):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"

    # --- Γράφημα 1: ημερήσια κατανομή (διακύμανση / συγκέντρωση ενδιαφέροντος) ---
    bar = BarChart()
    bar.type = "col"
    bar.title = "Νέα links ανά ημέρα — Δασικές Πυρκαγιές Ελλάδα 2026"
    bar.y_axis.title = "Αριθμός links"
    bar.x_axis.title = "Ημερομηνία"
    bar.height, bar.width = 9.5, 34
    bar.gapWidth = 40
    bar.legend = None
    data = Reference(ws, min_col=2, min_row=1, max_row=last_data_row)
    cats = Reference(ws, min_col=1, min_row=first_data_row, max_row=last_data_row)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.x_axis.tickLblSkip = 3
    bar.x_axis.tickMarkSkip = 3
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    ws.add_chart(bar, "E2")

    # --- Γράφημα 2: σωρευτική πορεία συλλογής ---
    line = LineChart()
    line.title = "Σωρευτικό σύνολο συνδέσμων"
    line.y_axis.title = "Σύνολο links"
    line.x_axis.title = "Ημερομηνία"
    line.height, line.width = 9.5, 34
    line.legend = None
    cum = Reference(ws, min_col=3, min_row=1, max_row=last_data_row)
    line.add_data(cum, titles_from_data=True)
    line.set_categories(cats)
    line.x_axis.tickLblSkip = 3
    line.x_axis.tickMarkSkip = 3
    line.x_axis.delete = False
    line.y_axis.delete = False
    ws.add_chart(line, "E22")

    # --- Φύλλο σύνοψης ανά μήνα ---
    ws2 = wb.create_sheet("Σύνοψη ανά μήνα")
    ws2.append(["Μήνας", "Links", "Ημέρες με links", "Μέγιστο ημέρας", "Ημερομηνία μεγίστου"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    by_month = collections.defaultdict(list)
    for d, n in per_day.items():
        if d <= last_day:
            by_month[(d.year, d.month)].append((d, n))
    for (y, m) in sorted(by_month):
        items = by_month[(y, m)]
        peak_day, peak_n = max(items, key=lambda t: (t[1], t[0]))
        ws2.append([f"{MONTHS[m - 1]} {y}", sum(n for _, n in items), len(items),
                    peak_n, peak_day.strftime("%d/%m")])
    ws2.append(["Μόνιμες / ζωντανές πηγές", undated, "—", "—", "—"])
    ws2.append(["ΣΥΝΟΛΟ", sum(n for _, n in per_day.items() if _ <= last_day) + undated,
                "—", "—", "—"])
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=5):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
    ws2[f"A{ws2.max_row}"].font = Font(bold=True)
    ws2[f"B{ws2.max_row}"].font = Font(bold=True)
    for i, w in enumerate([26, 12, 16, 16, 20], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT)
    return last_data_row - first_data_row + 1, running


def main():
    last_day = (dt.date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1
                else dt.date.today())
    per_day, undated = parse_links()
    if per_day:
        last_day = max(last_day, max(per_day))
    days, total = build(per_day, undated, last_day)
    print(f"OK: {days} ημέρες ({SEASON_START:%d/%m} → {last_day:%d/%m}), "
          f"σύνολο {total} links ({total - undated} με ημερομηνία + {undated} μόνιμες)")


if __name__ == "__main__":
    main()
